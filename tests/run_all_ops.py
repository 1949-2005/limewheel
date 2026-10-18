"""Smoke test: build sample files and run every non-interactive operation on them.

    python tests/run_all_ops.py [workdir]
"""
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from PySide6.QtWidgets import QApplication  # noqa: E402

app = QApplication([])

from limewheel.app import Ctx, execute  # noqa: E402
from limewheel.ops import MODES, ops_for  # noqa: E402
from limewheel.ops.media import ff  # noqa: E402
from limewheel import settings as settings_mod  # noqa: E402

work = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp(prefix="limewheel-test-"))
if work.exists():
    shutil.rmtree(work)
work.mkdir(parents=True)
src = work / "src"
src.mkdir()


def make_samples():
    from PIL import Image, ImageDraw
    im = Image.new("RGBA", (320, 200), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.ellipse((20, 20, 180, 180), fill=(255, 138, 0, 255))
    d.rectangle((200, 40, 300, 160), fill=(40, 120, 220, 200))
    im.save(src / "sample.png")
    im.convert("RGB").save(src / "photo.jpg", quality=90)
    frames = [Image.new("RGB", (120, 80), c) for c in ("red", "green", "blue")]
    frames[0].save(src / "anim.gif", save_all=True, append_images=frames[1:], duration=200, loop=0)
    (src / "logo.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="64" height="64">'
                                  '<circle cx="32" cy="32" r="28" fill="orange"/></svg>')
    ff("-f", "lavfi", "-i", "testsrc=size=320x240:rate=25:duration=3", "-f", "lavfi",
       "-i", "sine=frequency=440:duration=3", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac",
       "-shortest", src / "clip.mp4")
    ff("-f", "lavfi", "-i", "sine=frequency=660:duration=3", src / "tone.wav")
    import fitz
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_text((72, 72), f"Hello from page {i + 1}", fontsize=18)
        page.insert_image(fitz.Rect(72, 100, 272, 225), filename=str(src / "photo.jpg"))
    doc.save(src / "doc.pdf")
    from docx import Document
    dx = Document()
    dx.add_heading("Report", 1)
    dx.add_paragraph("Some **body** text.")
    dx.save(src / "report.docx")
    (src / "notes.md").write_text("# Title\n\n- one\n- **two**\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    (src / "plain.txt").write_text("Line one\n\nSecond paragraph.\n")
    (src / "page.html").write_text("<h1>Hi</h1><p>Paragraph <b>bold</b></p>")
    (src / "table.csv").write_text("name,qty\napple,3\npear,5\n")
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["x", "y"])
    wb.active.append([1, 2])
    wb.save(src / "book.xlsx")
    (src / "subs.srt").write_text("1\n00:00:01,000 --> 00:00:02,500\nHello\n\n2\n00:00:03,000 --> 00:00:04,000\nWorld\n")
    folder = src / "folder"
    folder.mkdir()
    (folder / "a.txt").write_text("a")
    import zipfile
    with zipfile.ZipFile(src / "bundle.zip", "w") as z:
        z.writestr("bundle/inner.txt", "inner")


make_samples()


class FakeApp:
    settings = settings_mod.load()

    def float_files(self, paths):
        pass


ctx = Ctx.__new__(Ctx)
ctx.app, ctx.settings = FakeApp(), FakeApp.settings

SKIP = {"act.copy", "act.move", "act.terminal", "act.reveal"}  # need dialogs or open windows
results = []
samples = sorted(p for p in src.iterdir())
groups = [[p] for p in samples] + [[src / "sample.png", src / "photo.jpg"], [src / "doc.pdf", src / "doc.pdf"],
                                    [src / "clip.mp4", src / "clip.mp4"], [src / "tone.wav", src / "tone.wav"]]
for paths in groups:
    for mode in MODES:
        for op in ops_for(paths, mode):
            if op.prompt or op.id in SKIP:
                continue
            if len(paths) > 1 and op.min_files < 2:
                continue
            case = work / "out" / f"{paths[0].stem}-{len(paths)}-{op.id}"
            case.mkdir(parents=True)
            local = []
            for p in paths:
                dst = case / p.name
                if not dst.exists():
                    (shutil.copytree if p.is_dir() else shutil.copy2)(p, dst)
                local.append(dst)
            if len(paths) == 2 and paths[0] == paths[1]:
                second = case / f"{paths[0].stem}-b{paths[0].suffix}"
                shutil.copy2(paths[0], second)
                local = [local[0], second]
            try:
                outs = execute(op, local, {}, ctx)
                ok = all(Path(o).exists() and (Path(o).is_dir() or Path(o).stat().st_size > 0) for o in outs)
                results.append((ok, op.id, paths[0].name, ", ".join(Path(o).name for o in outs)))
            except Exception as e:
                results.append((False, op.id, paths[0].name, f"{type(e).__name__}: {e}"))

# Prompted tools with fixed options
from limewheel.ops.base import REGISTRY  # noqa: E402
byid = {o.id: o for o in REGISTRY}
for oid, p, opts in [("img.resize", "sample.png", {"pct": 50}),
                     ("vid.trim", "clip.mp4", {"start": 0.5, "end": 2.0}),
                     ("vid.split", "clip.mp4", {"parts": 2}),
                     ("aud.trim", "tone.wav", {"start": 1.0, "end": None}),
                     ("aud.split", "tone.wav", {"parts": 3})]:
    case = work / "out" / f"prompt-{oid}"
    case.mkdir(parents=True)
    dst = case / p
    shutil.copy2(src / p, dst)
    try:
        outs = execute(byid[oid], [dst], opts, ctx)
        results.append((all(Path(o).exists() for o in outs), oid, p, ", ".join(Path(o).name for o in outs)))
    except Exception as e:
        results.append((False, oid, p, f"{type(e).__name__}: {e}"))

fails = [r for r in results if not r[0]]
for ok, oid, name, info in results:
    print(f"{'OK  ' if ok else 'FAIL'} {oid:20} {name:14} {info[:110]}")
print(f"\n{len(results) - len(fails)}/{len(results)} passed  (output in {work})")
sys.exit(1 if fails else 0)
