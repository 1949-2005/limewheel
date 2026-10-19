"""PDF, document, spreadsheet, slide and subtitle conversions."""
from __future__ import annotations

import csv
import html as htmlmod
import json
import re
import shutil
import tempfile
from functools import cache
from html.parser import HTMLParser
from pathlib import Path

from .base import OpError, ext_of, out_dir, out_path, register, run_proc

# ------------------------------------------------------------ office backends

WORD_IN = frozenset("doc docx odt rtf".split())
SHEET_OFFICE_IN = frozenset("xls xlsx ods".split())
SLIDE_IN = frozenset("ppt pptx odp".split())


def _has_progid(progid: str) -> bool:
    import winreg
    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, progid))
        return True
    except OSError:
        return False


@cache
def soffice() -> str | None:
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if Path(c).exists():
            return c
    return shutil.which("soffice")


@cache
def ms_office(app: str) -> bool:
    return _has_progid(f"{app}.Application")


def can_office(kind: str) -> bool:
    app = {"word": "Word", "sheet": "Excel", "slide": "PowerPoint"}[kind]
    return bool(soffice()) or ms_office(app)


def office_convert(src: Path, ext: str, kind: str) -> Path:
    dst = out_path(src, ext)
    app = {"word": "Word", "sheet": "Excel", "slide": "PowerPoint"}[kind]
    if ms_office(app):
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        try:
            o = win32com.client.DispatchEx(f"{app}.Application")
            try:
                if app == "Word":
                    o.Visible = False
                    d = o.Documents.Open(str(src), ReadOnly=True)
                    d.SaveAs2(str(dst), FileFormat={"pdf": 17, "docx": 16}[ext])
                    d.Close(False)
                elif app == "Excel":
                    o.DisplayAlerts = False
                    wb = o.Workbooks.Open(str(src), ReadOnly=True)
                    wb.ExportAsFixedFormat(0, str(dst))
                    wb.Close(False)
                else:
                    pr = o.Presentations.Open(str(src), ReadOnly=True, WithWindow=False)
                    pr.SaveAs(str(dst), 32)
                    pr.Close()
            finally:
                o.Quit()
        finally:
            pythoncom.CoUninitialize()
        return dst
    with tempfile.TemporaryDirectory() as tmp:
        run_proc([soffice(), "--headless", "--norestore", "--convert-to", ext, "--outdir", tmp, src], timeout=300)
        produced = next(Path(tmp).glob(f"*.{ext}"), None)
        if not produced:
            raise OpError("LibreOffice produced no output")
        shutil.move(produced, dst)
    return dst


# ---------------------------------------------------------------- html utils

class _Text(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "table"}

    def __init__(self):
        super().__init__()
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip -= 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_to_text(html: str) -> str:
    p = _Text()
    p.feed(html)
    return re.sub(r"\n{3,}", "\n\n", "".join(p.out)).strip() + "\n"


CSS = ("body{font-family:'Segoe UI',Arial;font-size:11pt;line-height:1.4}"
       "table{border-collapse:collapse}td,th{border:1px solid #999;padding:3px 6px}"
       "pre,code{font-family:Consolas,monospace;background:#f4f4f4}h1,h2,h3{color:#222}")


def html_to_pdf(html: str, dst: Path, base: Path | None = None, landscape=False) -> Path:
    from PySide6.QtCore import QMarginsF, QUrl
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument

    w = QPdfWriter(str(dst))
    w.setPageSize(QPageSize(QPageSize.A4))
    w.setPageOrientation(QPageLayout.Landscape if landscape else QPageLayout.Portrait)
    w.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
    doc = QTextDocument()
    if base:
        doc.setBaseUrl(QUrl.fromLocalFile(str(base) + "/"))
    doc.setDefaultStyleSheet(CSS)
    doc.setHtml(html)
    doc.print_(w)
    return dst


def wrap_html(body: str, title: str = "") -> str:
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{htmlmod.escape(title)}</title>"
            f"<style>{CSS}</style></head><body>{body}</body></html>")


def md_to_html(text: str) -> str:
    import markdown
    return markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])


def text_to_html(text: str) -> str:
    return "".join(f"<p>{htmlmod.escape(par).replace(chr(10), '<br>')}</p>"
                   for par in re.split(r"\n\s*\n", text) if par.strip())


def read_text(p: Path) -> str:
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            if enc == "utf-16" and not raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def md_to_docx(text: str, dst: Path) -> Path:
    from docx import Document

    doc = Document()

    def add_runs(par, line):
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", line):
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                par.add_run(part[1:-1]).font.name = "Consolas"
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                par.add_run(part[1:-1]).italic = True
            elif part:
                par.add_run(part)

    buf, in_code = [], False
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            par = doc.add_paragraph()
            if line:
                par.add_run(line).font.name = "Consolas"
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            doc.add_heading(m[2].strip(), level=min(len(m[1]), 4))
        elif re.match(r"^\s*[-*+]\s+", line):
            add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*+]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        elif line.strip():
            add_runs(doc.add_paragraph(), line.strip())
    doc.save(dst)
    return dst


# ------------------------------------------------------------------------ PDF

def _fitz():
    import fitz
    return fitz


def _pdf_images(ext):
    def run(p: Path, opts, ctx):
        doc = _fitz().open(p)
        try:
            if doc.page_count == 1:
                dst = out_path(p, ext)
                doc[0].get_pixmap(dpi=300).save(dst)
                return dst
            folder = out_dir(p, "-pages")
            for i, page in enumerate(doc, 1):
                page.get_pixmap(dpi=300).save(folder / f"{p.stem}-p{i:03d}.{ext}")
            return folder
        finally:
            doc.close()
    return run


def _pdf_text(p: Path) -> list[str]:
    doc = _fitz().open(p)
    try:
        return [page.get_text() for page in doc]
    finally:
        doc.close()


def _pdf_to_txt(p, opts, ctx):
    dst = out_path(p, "txt")
    dst.write_text("\n\f\n".join(_pdf_text(p)), encoding="utf-8")
    return dst


def _pdf_to_md(p, opts, ctx):
    try:
        import pymupdf4llm
        md = pymupdf4llm.to_markdown(str(p))
    except ImportError:
        pages = _pdf_text(p)
        md = "\n\n---\n\n".join(re.sub(r"[ \t]+\n", "\n", t).strip() for t in pages)
    dst = out_path(p, "md")
    dst.write_text(md, encoding="utf-8")
    return dst


def _pdf_to_html(p, opts, ctx):
    doc = _fitz().open(p)
    try:
        body = "".join(f"<section>{page.get_text('xhtml')}</section><hr>" for page in doc)
    finally:
        doc.close()
    dst = out_path(p, "html")
    dst.write_text(wrap_html(body, p.stem), encoding="utf-8")
    return dst


def _pdf_to_docx(p, opts, ctx):
    """Rebuild a PDF as an editable Word file: headings, bold/italic runs, colours, tables and images."""
    import io
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt, RGBColor, Emu

    fitz = _fitz()
    pdf = fitz.open(p)
    doc = Document()
    sec = doc.sections[0]
    try:
        r0 = pdf[0].rect
        sec.page_width, sec.page_height = Emu(int(r0.width * 12700)), Emu(int(r0.height * 12700))
    except Exception:
        pass
    usable = sec.page_width - sec.left_margin - sec.right_margin
    # body font size = most common span size in the document
    sizes: dict[float, int] = {}
    for page in pdf:
        for b in page.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                for s in l["spans"]:
                    k = round(s["size"], 1)
                    sizes[k] = sizes.get(k, 0) + len(s["text"])
    body = max(sizes, key=sizes.get) if sizes else 11.0

    for pi, page in enumerate(pdf):
        if pi:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        tables = []
        try:
            tables = page.find_tables().tables
        except Exception:
            pass
        tboxes = [fitz.Rect(t.bbox) for t in tables]
        items = []                               # (y, kind, payload)
        for t, box in zip(tables, tboxes):
            items.append((box.y0, "table", t))
        for b in page.get_text("dict")["blocks"]:
            rect = fitz.Rect(b["bbox"])
            if any(rect.intersects(tb) and (rect & tb).get_area() > 0.6 * rect.get_area() for tb in tboxes):
                continue
            items.append((rect.y0, "image" if b["type"] == 1 else "text", b))
        for _, kind, b in sorted(items, key=lambda x: x[0]):
            if kind == "table":
                rows = [[("" if c is None else str(c)) for c in row] for row in b.extract()]
                if not rows:
                    continue
                tbl = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                tbl.style = "Table Grid"
                for i, row in enumerate(rows):
                    for j, c in enumerate(row):
                        tbl.cell(i, j).text = c
                continue
            if kind == "image":
                try:
                    w = min(usable, Emu(int((b["bbox"][2] - b["bbox"][0]) * 12700)))
                    doc.add_picture(io.BytesIO(b["image"]), width=w)
                except Exception:
                    pass
                continue
            spans = [s for l in b["lines"] for s in l["spans"] if s["text"].strip()]
            if not spans:
                continue
            size = max(s["size"] for s in spans)
            level = 1 if size >= body * 1.6 else 2 if size >= body * 1.3 else 3 if size >= body * 1.12 else 0
            par = doc.add_heading(level=level) if level else doc.add_paragraph()
            for li, l in enumerate(b["lines"]):
                for s in l["spans"]:
                    text = s["text"]
                    if not text:
                        continue
                    run = par.add_run(text)
                    run.bold = bool(s["flags"] & 16) or None
                    run.italic = bool(s["flags"] & 2) or None
                    if not level:
                        run.font.size = Pt(round(s["size"] * 2) / 2)
                    col = s.get("color", 0)
                    if col:
                        run.font.color.rgb = RGBColor((col >> 16) & 255, (col >> 8) & 255, col & 255)
                if li < len(b["lines"]) - 1:
                    par.add_run(" ")
    pdf.close()
    dst = out_path(p, "docx")
    doc.save(dst)
    return dst


def _pdf_to_pptx(p, opts, ctx):
    import io
    from pptx import Presentation
    from pptx.util import Emu

    doc = _fitz().open(p)
    try:
        prs = Presentation()
        r = doc[0].rect
        prs.slide_width = Emu(int(r.width * 12700))
        prs.slide_height = Emu(int(r.height * 12700))
        blank = prs.slide_layouts[6]
        for page in doc:
            slide = prs.slides.add_slide(blank)
            img = io.BytesIO(page.get_pixmap(dpi=150).tobytes("png"))
            slide.shapes.add_picture(img, 0, 0, prs.slide_width, prs.slide_height)
    finally:
        doc.close()
    dst = out_path(p, "pptx")
    prs.save(dst)
    return dst


def _pdf_to_xlsx(p, opts, ctx):
    from openpyxl import Workbook

    doc = _fitz().open(p)
    wb = Workbook()
    wb.remove(wb.active)
    try:
        for i, page in enumerate(doc, 1):
            for j, table in enumerate(page.find_tables().tables, 1):
                ws = wb.create_sheet(f"p{i}-t{j}"[:31])
                for row in table.extract():
                    ws.append([("" if c is None else c) for c in row])
    finally:
        doc.close()
    if not wb.worksheets:
        raise OpError("No tables were found in this PDF")
    dst = out_path(p, "xlsx")
    wb.save(dst)
    return dst


PDF = frozenset({"pdf"})
register("pdf.docx", "to DOCX", "convert", PDF, _pdf_to_docx, badge="DOCX")
register("pdf.jpg", "to JPG", "convert", PDF, _pdf_images("jpg"), badge="JPG")
register("pdf.png", "to PNG", "convert", PDF, _pdf_images("png"), badge="PNG")
register("pdf.txt", "to TXT", "convert", PDF, _pdf_to_txt, badge="TXT")
register("pdf.md", "to Markdown", "convert", PDF, _pdf_to_md, badge="MD")
register("pdf.html", "to HTML", "convert", PDF, _pdf_to_html, badge="HTML")
register("pdf.pptx", "to PPTX", "convert", PDF, _pdf_to_pptx, badge="PPTX")
register("pdf.xlsx", "Tables → XLSX", "convert", PDF, _pdf_to_xlsx, badge="XLSX")


def _pdf_merge(paths, opts, ctx):
    fitz = _fitz()
    out = fitz.open()
    for p in paths:
        with fitz.open(p) as d:
            out.insert_pdf(d)
    dst = out_path(paths[0], "pdf", "-merged")
    out.save(dst, garbage=3, deflate=True)
    out.close()
    return dst


def _pdf_split(p, opts, ctx):
    fitz = _fitz()
    folder = out_dir(p, "-split")
    with fitz.open(p) as d:
        for i in range(d.page_count):
            one = fitz.open()
            one.insert_pdf(d, from_page=i, to_page=i)
            one.save(folder / f"{p.stem}-p{i + 1:03d}.pdf")
            one.close()
    return folder


def _pdf_compress(p, opts, ctx):
    fitz = _fitz()
    dst = out_path(p, "pdf", "-compressed")
    with fitz.open(p) as d:
        if hasattr(d, "rewrite_images"):
            try:
                d.rewrite_images(dpi_threshold=160, dpi_target=120, quality=70)
            except Exception:
                pass
        d.save(dst, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
    return dst


def _pdf_rotate(p, opts, ctx):
    fitz = _fitz()
    dst = out_path(p, "pdf", "-rotated")
    with fitz.open(p) as d:
        for page in d:
            page.set_rotation((page.rotation + 90) % 360)
        d.save(dst, garbage=3, deflate=True)
    return dst


def _pdf_extract_images(p, opts, ctx):
    fitz = _fitz()
    folder = out_dir(p, "-images")
    n = 0
    with fitz.open(p) as d:
        seen = set()
        for page in d:
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                info = d.extract_image(xref)
                n += 1
                (folder / f"{p.stem}-img{n:03d}.{info['ext']}").write_bytes(info["image"])
    if not n:
        folder.rmdir()
        raise OpError("This PDF has no embedded images")
    return folder


register("pdf.merge", "Merge", "tools", PDF, _pdf_merge, batch=True, min_files=2)
register("pdf.split", "Split pages", "tools", PDF, _pdf_split)
register("pdf.compress", "Compress", "tools", PDF, _pdf_compress)
register("pdf.rotate", "Rotate 90°", "tools", PDF, _pdf_rotate)
register("pdf.images", "Extract images", "tools", PDF, _pdf_extract_images)


# ------------------------------------------------------------- word documents

def _docx_html(p: Path) -> str:
    import mammoth
    with open(p, "rb") as f:
        return mammoth.convert_to_html(f).value


def _docx_to(kind):
    def run(p: Path, opts, ctx):
        import mammoth
        dst = out_path(p, kind)
        if kind == "html":
            dst.write_text(wrap_html(_docx_html(p), p.stem), encoding="utf-8")
        elif kind == "md":
            with open(p, "rb") as f:
                dst.write_text(mammoth.convert_to_markdown(f).value, encoding="utf-8")
        else:
            with open(p, "rb") as f:
                dst.write_text(mammoth.extract_raw_text(f).value, encoding="utf-8")
        return dst
    return run


DOCX = frozenset({"docx"})
register("docx.html", "to HTML", "convert", DOCX, _docx_to("html"), badge="HTML")
register("docx.md", "to Markdown", "convert", DOCX, _docx_to("md"), badge="MD")
register("docx.txt", "to TXT", "convert", DOCX, _docx_to("txt"), badge="TXT")
register("docx.pdf.basic", "to PDF", "convert", DOCX,
         lambda p, o, c: html_to_pdf(_docx_html(p), out_path(p, "pdf"), p.parent),
         badge="PDF", main_thread=True, available=lambda: not can_office("word"))
register("office.word.pdf", "to PDF", "convert", WORD_IN,
         lambda p, o, c: office_convert(p, "pdf", "word"), badge="PDF", available=lambda: can_office("word"))
register("office.word.docx", "to DOCX", "convert", WORD_IN - {"docx"},
         lambda p, o, c: office_convert(p, "docx", "word"), badge="DOCX", available=lambda: can_office("word"))
register("office.sheet.pdf", "to PDF", "convert", SHEET_OFFICE_IN,
         lambda p, o, c: office_convert(p, "pdf", "sheet"), badge="PDF", available=lambda: can_office("sheet"))
register("office.slide.pdf", "to PDF", "convert", SLIDE_IN,
         lambda p, o, c: office_convert(p, "pdf", "slide"), badge="PDF", available=lambda: can_office("slide"))


# ------------------------------------------------------- markdown / txt / html

MD, TXT, HTML = frozenset({"md"}), frozenset({"txt"}), frozenset({"html"})


def _write(p, ext, content):
    dst = out_path(p, ext)
    dst.write_text(content, encoding="utf-8")
    return dst


register("md.html", "to HTML", "convert", MD,
         lambda p, o, c: _write(p, "html", wrap_html(md_to_html(read_text(p)), p.stem)), badge="HTML")
register("md.pdf", "to PDF", "convert", MD,
         lambda p, o, c: html_to_pdf(md_to_html(read_text(p)), out_path(p, "pdf"), p.parent),
         badge="PDF", main_thread=True)
register("md.docx", "to DOCX", "convert", MD,
         lambda p, o, c: md_to_docx(read_text(p), out_path(p, "docx")), badge="DOCX")
register("md.txt", "to TXT", "convert", MD,
         lambda p, o, c: _write(p, "txt", html_to_text(md_to_html(read_text(p)))), badge="TXT")

register("txt.pdf", "to PDF", "convert", TXT,
         lambda p, o, c: html_to_pdf(text_to_html(read_text(p)), out_path(p, "pdf")),
         badge="PDF", main_thread=True)
register("txt.docx", "to DOCX", "convert", TXT,
         lambda p, o, c: md_to_docx(read_text(p), out_path(p, "docx")), badge="DOCX")
register("txt.md", "to Markdown", "convert", TXT,
         lambda p, o, c: _write(p, "md", read_text(p)), badge="MD")
register("txt.html", "to HTML", "convert", TXT,
         lambda p, o, c: _write(p, "html", wrap_html(text_to_html(read_text(p)), p.stem)), badge="HTML")

register("html.pdf", "to PDF", "convert", HTML,
         lambda p, o, c: html_to_pdf(read_text(p), out_path(p, "pdf"), p.parent), badge="PDF", main_thread=True)
register("html.txt", "to TXT", "convert", HTML,
         lambda p, o, c: _write(p, "txt", html_to_text(read_text(p))), badge="TXT")
register("html.docx", "to DOCX", "convert", HTML,
         lambda p, o, c: md_to_docx(html_to_text(read_text(p)), out_path(p, "docx")), badge="DOCX")


# --------------------------------------------------------------- spreadsheets

SHEETS = frozenset("xlsx xls csv tsv ods".split())


def _read_ods(p: Path) -> list[tuple[str, list[list]]]:
    from odf import teletype
    from odf.opendocument import load
    from odf.table import Table, TableCell, TableRow
    from odf.text import P
    out = []
    for t in load(str(p)).spreadsheet.getElementsByType(Table):
        rows = []
        for tr in t.getElementsByType(TableRow):
            row = []
            for tc in tr.getElementsByType(TableCell):
                rep = min(int(tc.getAttribute("numbercolumnsrepeated") or 1), 512)
                vtype, val = tc.getAttribute("valuetype"), tc.getAttribute("value")
                text = "\n".join(teletype.extractText(x) for x in tc.getElementsByType(P))
                v = (float(val) if vtype in ("float", "percentage", "currency") and val else text)
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                row += [v] * rep
            while row and row[-1] in ("", None):
                row.pop()
            rows += [row] * min(int(tr.getAttribute("numberrowsrepeated") or 1), 512)
        while rows and not rows[-1]:
            rows.pop()
        out.append((t.getAttribute("name") or "Sheet", rows))
    return out


def _write_ods(sheets, dst: Path) -> Path:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P
    doc = OpenDocumentSpreadsheet()
    for name, rows in sheets:
        t = Table(name=str(name)[:31] or "Sheet")
        for row in rows:
            tr = TableRow()
            for v in row:
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    tc = TableCell(valuetype="float", value=v)
                else:
                    tc = TableCell(valuetype="string")
                tc.addElement(P(text="" if v is None else str(v)))
                tr.addElement(tc)
            t.addElement(tr)
        doc.spreadsheet.addElement(t)
    doc.save(str(dst))
    return dst


def read_sheets(p: Path) -> list[tuple[str, list[list]]]:
    e = ext_of(p)
    if e == "ods":
        return _read_ods(p)
    if e == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            return [(ws.title, [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)])
                    for ws in wb.worksheets]
        finally:
            wb.close()
    if e == "xls":
        import xlrd
        book = xlrd.open_workbook(p)
        return [(s.name, [s.row_values(r) for r in range(s.nrows)]) for s in book.sheets()]
    text = read_text(p)
    delim = "\t" if e == "tsv" else None
    if delim is None:
        try:
            delim = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ","
    return [(p.stem, list(csv.reader(text.splitlines(), delimiter=delim)))]


def _sheet_to(ext):
    def run(p: Path, opts, ctx):
        sheets = read_sheets(p)
        if ext == "ods":
            return _write_ods(sheets, out_path(p, "ods"))
        if ext == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            wb.remove(wb.active)
            for name, rows in sheets:
                ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet")
                for row in rows:
                    ws.append(row)
            dst = out_path(p, "xlsx")
            wb.save(dst)
            return dst
        if ext in ("csv", "tsv"):
            outs = []
            for name, rows in sheets:
                suffix = f"-{name}" if len(sheets) > 1 else ""
                dst = out_path(p, ext, suffix)
                with open(dst, "w", newline="", encoding="utf-8-sig") as f:
                    csv.writer(f, delimiter="\t" if ext == "tsv" else ",").writerows(rows)
                outs.append(dst)
            return outs
        if ext == "json":
            data = {}
            for name, rows in sheets:
                if rows:
                    head = [str(h) for h in rows[0]]
                    data[name] = [dict(zip(head, r)) for r in rows[1:]]
            return _write(p, "json", json.dumps(data if len(data) != 1 else next(iter(data.values())),
                                                indent=2, ensure_ascii=False, default=str))
        body = ""
        for name, rows in sheets:
            body += f"<h2>{htmlmod.escape(str(name))}</h2><table>"
            for i, row in enumerate(rows):
                tag = "th" if i == 0 else "td"
                body += "<tr>" + "".join(f"<{tag}>{htmlmod.escape(str(c))}</{tag}>" for c in row) + "</tr>"
            body += "</table>"
        if ext == "html":
            return _write(p, "html", wrap_html(body, p.stem))
        return html_to_pdf(body, out_path(p, "pdf"), landscape=True)
    return run


for _ext in ("xlsx", "ods", "csv", "tsv", "json", "html"):
    register(f"sheet.{_ext}", f"to {_ext.upper()}", "convert", SHEETS, _sheet_to(_ext),
             badge=_ext.upper(), skip_same=_ext)
register("sheet.pdf", "to PDF", "convert", SHEETS, _sheet_to("pdf"), badge="PDF", main_thread=True,
         available=lambda: not can_office("sheet"))
register("sheet.pdf.csv", "to PDF", "convert", frozenset({"csv", "tsv"}), _sheet_to("pdf"), badge="PDF",
         main_thread=True, available=lambda: can_office("sheet"))


# ------------------------------------------------------------------ subtitles

SUBS = frozenset({"srt", "vtt"})
_TS = re.compile(r"(\d+:)?(\d{1,2}):(\d{2})[.,](\d{3})")


def _cues(p: Path):
    cues = []
    for block in re.split(r"\n\s*\n", read_text(p).replace("\r\n", "\n")):
        lines = [l for l in block.strip().splitlines() if l.strip()]
        for i, line in enumerate(lines):
            if "-->" in line:
                a, b = line.split("-->")[:2]
                cues.append((_TS.search(a)[0], _TS.search(b)[0], "\n".join(lines[i + 1:])))
                break
    return cues


def _ts(t, sep):
    m = _TS.match(t)
    h = (m[1] or "0:")[:-1]
    return f"{int(h):02d}:{int(m[2]):02d}:{m[3]}{sep}{m[4]}"


def _sub_to(ext):
    def run(p: Path, opts, ctx):
        cues = _cues(p)
        if ext == "vtt":
            body = "WEBVTT\n\n" + "\n".join(f"{_ts(a, '.')} --> {_ts(b, '.')}\n{t}\n" for a, b, t in cues)
        elif ext == "srt":
            body = "\n".join(f"{i}\n{_ts(a, ',')} --> {_ts(b, ',')}\n{t}\n" for i, (a, b, t) in enumerate(cues, 1))
        else:
            body = "\n".join(re.sub(r"<[^>]+>", "", t) for _, _, t in cues) + "\n"
        return _write(p, ext, body)
    return run


for _ext in ("srt", "vtt", "txt"):
    register(f"sub.{_ext}", f"to {_ext.upper()}", "convert", SUBS, _sub_to(_ext), badge=_ext.upper(),
             skip_same=_ext)
